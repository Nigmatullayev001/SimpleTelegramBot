import logging
import tempfile
from aiogram import Bot, Dispatcher, types
from aiogram.filters import Command
from aiogram.types import FSInputFile
from aiogram.enums import ParseMode
from aiogram.utils.markdown import hbold
from pptx import Presentation
from pptx.util import Inches, Pt
import aiohttp
import asyncio
import os

API_TOKEN = "7389652173:AAEjyKdU2KXMteLCUOUNdqDCf5b7iFcbiiU"
logging.basicConfig(level=logging.INFO)

bot = Bot(token=API_TOKEN)
dp = Dispatcher()

# Store user session slide content
user_sessions = {}

@dp.message(Command("start"))
async def cmd_start(message: types.Message):
    await message.answer("Send me texts and images for your slides. When done, send /done to get the PowerPoint file.")

@dp.message(Command("done"))
async def cmd_done(message: types.Message):
    user_id = message.from_user.id
    slides = user_sessions.get(user_id, [])
    if not slides:
        await message.answer("You haven't sent any content for slides.")
        return

    # Create PPTX
    prs = Presentation()
    temp_files = []

    for item in slides:
        slide_layout = prs.slide_layouts[5]
        slide = prs.slides.add_slide(slide_layout)

        if item['type'] == 'text':
            txBox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(5))
            tf = txBox.text_frame
            p = tf.add_paragraph()
            p.text = item['data']
            p.font.size = Pt(32)
        elif item['type'] == 'image':
            with tempfile.NamedTemporaryFile(delete=False, suffix=".jpg") as tmp:
                tmp.write(item['data'])
                img_path = tmp.name
                temp_files.append(img_path)
            slide.shapes.add_picture(img_path, Inches(1), Inches(1), width=Inches(7))

    # Save PPTX
    with tempfile.NamedTemporaryFile(delete=False, suffix=".pptx") as tmp:
        prs.save(tmp.name)
        pptx_path = tmp.name

    file = FSInputFile(pptx_path, filename="slides.pptx")
    await message.answer_document(file)

    # cleanup
    os.unlink(pptx_path)
    for f in temp_files:
        os.unlink(f)

    user_sessions[user_id] = []


@dp.message()
async def collect_content(message: types.Message):
    user_id = message.from_user.id
    user_sessions.setdefault(user_id, [])

    if message.photo:
        # Get largest photo
        photo = message.photo[-1]
        file = await bot.get_file(photo.file_id)
        file_path = file.file_path
        file_url = f"https://api.telegram.org/file/bot{API_TOKEN}/{file_path}"
        async with aiohttp.ClientSession() as session:
            async with session.get(file_url) as resp:
                img_bytes = await resp.read()
        user_sessions[user_id].append({'type': 'image', 'data': img_bytes})
        await message.answer("✅ Added image to your slides.")
    elif message.text:
        user_sessions[user_id].append({'type': 'text', 'data': message.text})
        await message.answer(f"📝 Added text: {hbold(message.text)}", parse_mode=ParseMode.HTML)
    else:
        await message.answer("Send text or images for slides. When done, send /done.")

if __name__ == "__main__":
    asyncio.run(dp.start_polling(bot))
