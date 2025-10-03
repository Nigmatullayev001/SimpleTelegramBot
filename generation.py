from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Image, Spacer, PageBreak
from reportlab.lib.enums import TA_CENTER
import os


def generate_file(fmt, uid, items, settings=None, template_name=None):
    """
    settings: {
        "text_size": int,
        "per_page": int
    }
    template_name: str -> 'green', 'blue', 'pink', ...
    """
    filename = f"results/{uid}_file.{fmt}"

    text_size = settings.get("text_size", 18) if settings else 18
    per_page = settings.get("per_page", 1) if settings else 1

    # Dinamik rasm o‘lchami per_page ga qarab
    if per_page == 1:
        img_w, img_h = 500, 380
    elif per_page == 2:
        img_w, img_h = 380, 280
    else:
        img_w, img_h = 250, 180

    if fmt == "pdf":
        doc = SimpleDocTemplate(filename, pagesize=A4)
        story = []

        styles = getSampleStyleSheet()
        styles.add(ParagraphStyle(name="Center", alignment=TA_CENTER, fontSize=text_size, leading=text_size+4))
        styles.add(ParagraphStyle(name="Caption", alignment=TA_CENTER, fontSize=int(text_size*0.8), leading=text_size, textColor="gray"))

        counter = 0
        for item in items:
            if item[0] == "text":
                story.append(Paragraph(item[1], styles["Center"]))
                story.append(Spacer(1, 15))
            elif item[0] == "image":
                if os.path.exists(item[1]):
                    img = Image(item[1], width=img_w, height=img_h)
                    img.hAlign = "CENTER"
                    story.append(img)
                    if item[2]:
                        story.append(Paragraph(item[2], styles["Caption"]))
                    story.append(Spacer(1, 20))

            counter += 1
            if counter >= per_page:
                story.append(PageBreak())
                counter = 0

        doc.build(story)
        return filename

    elif fmt == "docx":
        from docx import Document
        from docx.shared import Pt, Inches

        doc = Document()
        counter = 0

        for item in items:
            if item[0] == "text":
                para = doc.add_paragraph(item[1])
                run = para.runs[0]
                run.font.size = Pt(text_size)
            elif item[0] == "image":
                if os.path.exists(item[1]):
                    doc.add_picture(item[1], width=Inches(img_w / 100))
                    if item[2]:
                        para = doc.add_paragraph(item[2])
                        run = para.runs[0]
                        run.font.size = Pt(int(text_size * 0.8))

            counter += 1
            if counter >= per_page:
                doc.add_page_break()
                counter = 0

        doc.save(filename)
        return filename

    elif fmt == "pptx":
        return generate_ppt(uid, items, per_page, text_size, template_name or "default")


from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

def generate_ppt(uid, items, per_page, text_size, template_name):
    template_path = f"templates/{template_name}.pptx"
    if os.path.exists(template_path):
        prs = Presentation(template_path)
        # ✅ Default bo‘sh slaydni olib tashlaymiz
        if len(prs.slides) > 0:
            rId = prs.slides._sldIdLst[0].rId
            prs.part.drop_rel(rId)
            del prs.slides._sldIdLst[0]
    else:
        prs = Presentation()

    blank_slide_layout = prs.slide_layouts[6]

    # Rang tanlash
    if template_name == "green" or template_name == "black_green":
        text_color = RGBColor(0, 176, 80)
    elif template_name == "blue":
        text_color = RGBColor(0, 112, 192)
    elif template_name == "pink":
        text_color = RGBColor(192, 0, 192)
    else:
        text_color = RGBColor(0, 0, 0)

    # Dinamik rasm o‘lchami
    if per_page == 1:
        img_w, img_h = 6, 4.5
    elif per_page == 2:
        img_w, img_h = 4.5, 3
    else:
        img_w, img_h = 3.5, 2.5

    counter = 0
    slide = None

    for item in items:
        if counter == 0:  # yangi sahifa
            slide = prs.slides.add_slide(blank_slide_layout)

        if item[0] == "text":
            txBox = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(3))
            tf = txBox.text_frame
            p = tf.add_paragraph()
            p.text = item[1]
            p.font.size = Pt(text_size + 10)
            p.font.color.rgb = text_color
            p.alignment = 1

        elif item[0] == "image":
            if os.path.exists(item[1]):
                slide.shapes.add_picture(item[1], Inches(1.5), Inches(1), Inches(img_w), Inches(img_h))
                if item[2]:
                    txBox = slide.shapes.add_textbox(Inches(1), Inches(5.5), Inches(8), Inches(1))
                    tf = txBox.text_frame
                    p = tf.add_paragraph()
                    p.text = item[2]
                    p.font.size = Pt(int(text_size * 0.8))
                    p.font.color.rgb = text_color
                    p.alignment = 1

        counter += 1
        if counter >= per_page:
            counter = 0  # yangi sahifa uchun reset

    filename = f"results/{uid}_file.pptx"
    prs.save(filename)
    return filename

